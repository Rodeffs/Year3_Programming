from docx import Document


def main():
    document = Document("output.docx")

    table = document.tables[0]
    data = {}

    for row in range(1, 4):
        stat_name = table.cell(row, 0).text[1:]
        stat = table.cell(row, 2).text[1:]

        data[stat_name] = stat
    
    print("Данные по ATmega328:")
    print(data)


if __name__ == "__main__":
    main()
